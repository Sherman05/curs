"""
services/export.py — Экспорт отчётов в Excel (openpyxl).

Три отчёта:
  - export_group_performance(group_id)  — успеваемость группы (по предметам);
  - export_overall_report()             — общий отчёт (4 листа);
  - export_student_record(student_id)   — зачётная книжка студента.

Каждая функция возвращает bytes готового .xlsx (через BytesIO).
Цветовая заливка средних: <3.5 — красный, 3.5–4 — жёлтый, ≥4 — зелёный.
"""
import logging
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from models import db, Grade, Group, Student, Subject, Teacher, User
from services import analytics

logger = logging.getLogger(__name__)

# Стили.
HEADER_FONT = Font(bold=True)
FILL_RED = PatternFill("solid", fgColor="F8CBAD")     # < 3.5
FILL_YELLOW = PatternFill("solid", fgColor="FFE699")  # 3.5–4
FILL_GREEN = PatternFill("solid", fgColor="C6E0B4")   # ≥ 4


def _fill_for(avg) -> PatternFill | None:
    """Подобрать цвет заливки по среднему баллу."""
    if avg is None or avg == 0:
        return None
    if avg < 3.5:
        return FILL_RED
    if avg < 4:
        return FILL_YELLOW
    return FILL_GREEN


def _autosize(ws) -> None:
    """Авторазмер колонок по максимальной длине содержимого."""
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                col = cell.column_letter
                widths[col] = max(widths.get(col, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(w + 2, 50)


def _save(wb) -> bytes:
    """Сохранить книгу в байты."""
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# 1. Успеваемость группы
# ---------------------------------------------------------------------------
def export_group_performance(group_id: int) -> bytes:
    """Excel-отчёт по успеваемости группы (строка — студент, колонки — предметы)."""
    s = db.session
    group = s.get(Group, group_id)

    # Предметы, по которым у студентов группы есть оценки.
    subjects = (s.query(Subject.id, Subject.name)
                .join(Grade, Grade.subject_id == Subject.id)
                .join(Student, Student.id == Grade.student_id)
                .filter(Student.group_id == group_id)
                .distinct().order_by(Subject.name).all())

    wb = Workbook()
    ws = wb.active
    ws.title = "Успеваемость группы"

    # Шапка.
    header = ["ФИО", "Группа"] + [name for _, name in subjects] + ["Общий средний"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    students = (Student.query.filter_by(group_id=group_id)
                .join(User).order_by(User.full_name).all())

    for st in students:
        row = [st.user.full_name, group.name]
        # Средний по каждому предмету.
        for sid, _ in subjects:
            avg = s.query(func.avg(Grade.value)).filter(
                Grade.student_id == st.id, Grade.subject_id == sid).scalar()
            row.append(round(avg, 2) if avg is not None else "")
        overall = analytics.student_avg(s, st.id)
        row.append(overall)
        ws.append(row)

        # Заливка ячеек средних (по предметам + общий).
        r = ws.max_row
        for col_idx in range(3, 3 + len(subjects) + 1):
            cell = ws.cell(row=r, column=col_idx)
            if isinstance(cell.value, (int, float)):
                fill = _fill_for(cell.value)
                if fill:
                    cell.fill = fill

    _autosize(ws)
    logger.info("Сформирован Excel: успеваемость группы %d.", group_id)
    return _save(wb)


# ---------------------------------------------------------------------------
# 2. Общий отчёт (4 листа)
# ---------------------------------------------------------------------------
def export_overall_report() -> bytes:
    """Общий отчёт администратора: статистика, топ, группа риска, распределение."""
    s = db.session
    wb = Workbook()

    # Лист 1: общая статистика (ключ → значение).
    ws1 = wb.active
    ws1.title = "Общая статистика"
    ws1.append(["Показатель", "Значение"])
    for c in ws1[1]:
        c.font = HEADER_FONT
    stats = analytics.overall_stats(s)
    labels = {
        "total_students": "Студентов", "total_teachers": "Преподавателей",
        "total_groups": "Групп", "total_subjects": "Предметов",
        "overall_avg": "Общий средний балл", "at_risk_count": "В группе риска",
    }
    for key, label in labels.items():
        ws1.append([label, stats[key]])
    _autosize(ws1)

    # Лист 2: топ-10.
    ws2 = wb.create_sheet("Топ-10")
    ws2.append(["ФИО", "Группа", "Средний"])
    for c in ws2[1]:
        c.font = HEADER_FONT
    for r in analytics.top_students(s, limit=10):
        ws2.append([r["full_name"], r["group"], r["avg"]])
    _autosize(ws2)

    # Лист 3: группа риска.
    ws3 = wb.create_sheet("Группа риска")
    ws3.append(["ФИО", "Группа", "Средний"])
    for c in ws3[1]:
        c.font = HEADER_FONT
    for r in analytics.at_risk_students(s):
        ws3.append([r["full_name"], r["group"], r["avg"]])
        ws3.cell(row=ws3.max_row, column=3).fill = FILL_RED
    _autosize(ws3)

    # Лист 4: распределение оценок.
    ws4 = wb.create_sheet("Распределение оценок")
    ws4.append(["Оценка", "Количество"])
    for c in ws4[1]:
        c.font = HEADER_FONT
    dist = analytics.grade_distribution(s, scope="all")
    for grade in (2, 3, 4, 5):
        ws4.append([grade, dist[grade]])
    _autosize(ws4)

    logger.info("Сформирован Excel: общий отчёт.")
    return _save(wb)


# ---------------------------------------------------------------------------
# 3. Зачётная книжка студента
# ---------------------------------------------------------------------------
def export_student_record(student_id: int) -> bytes:
    """Excel-«зачётка»: все оценки + средние по предметам + общий средний."""
    s = db.session
    student = s.get(Student, student_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Зачётная книжка"

    # Все оценки (новые сверху), с преподавателем.
    ws.append(["Дата", "Предмет", "Балл", "Преподаватель", "Комментарий"])
    for c in ws[1]:
        c.font = HEADER_FONT

    rows = (s.query(Grade.date, Subject.name, Grade.value, User.full_name, Grade.comment)
            .join(Subject, Subject.id == Grade.subject_id)
            .outerjoin(Teacher, Teacher.id == Subject.teacher_id)
            .outerjoin(User, User.id == Teacher.user_id)
            .filter(Grade.student_id == student_id)
            .order_by(Grade.date.desc()).all())
    for date_, subj, value, teacher_name, comment in rows:
        ws.append([date_.strftime("%d.%m.%Y") if date_ else "", subj, value,
                   teacher_name or "—", comment or ""])

    # Пустая строка → блок средних по предметам.
    ws.append([])
    ws.append(["Средние по предметам"])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
    ws.append(["Предмет", "Средний", "Кол-во оценок"])
    for c in ws[ws.max_row]:
        c.font = HEADER_FONT
    for r in analytics.student_avg_by_subject(s, student_id):
        ws.append([r["subject"], r["avg"], r["count"]])
        ws.cell(row=ws.max_row, column=2).fill = _fill_for(r["avg"]) or PatternFill()

    # Общий средний.
    ws.append([])
    ws.append([f"Общий средний: {analytics.student_avg(s, student_id):.2f}"])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT

    _autosize(ws)
    logger.info("Сформирован Excel: зачётка студента %d.", student_id)
    return _save(wb)
